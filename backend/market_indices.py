"""Official monthly macro data and index levels. No synthetic fallbacks."""
import io
import logging
import re
import ssl
import time
from datetime import datetime
from functools import lru_cache
from pathlib import Path
from urllib.parse import urljoin, urlparse
from xml.etree import ElementTree
from zoneinfo import ZoneInfo
from zipfile import BadZipFile

import numpy as np
import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from requests.adapters import HTTPAdapter

from data_loader import DataError, _load_moex_prices, _load_yahoo_prices
from index_catalog import INDEX_BY_ID

logger = logging.getLogger('portfolio.indices')
ROSSTAT_URL = 'https://rosstat.gov.ru/statistics/price'
M2_URL = 'https://www.cbr.ru/vfs/statistics/credit_statistics/monetary_agg.xlsx'
FX_URL = 'https://www.cbr.ru/scripts/XML_dynamic.asp'
FX_IDS = {'USD': 'R01235', 'GBP': 'R01035', 'JPY': 'R01820'}
MONTHS = ['январь', 'февраль', 'март', 'апрель', 'май', 'июнь', 'июль', 'август', 'сентябрь', 'октябрь', 'ноябрь', 'декабрь']


class RosstatTLSAdapter(HTTPAdapter):
    """Additional official CA chain scoped exclusively to requests to Rosstat."""
    def init_poolmanager(self, *args, **kwargs):
        context = ssl.create_default_context(cafile=requests.certs.where())
        certs = Path(__file__).with_name('certificates')
        context.load_verify_locations(cafile=str(certs / 'russian_trusted_root.pem'))
        context.load_verify_locations(cafile=str(certs / 'russian_trusted_sub2024.pem'))
        kwargs['ssl_context'] = context
        return super().init_poolmanager(*args, **kwargs)


def _get(url, *, rosstat=False, params=None):
    try:
        with requests.Session() as session:
            if rosstat:
                session.mount('https://rosstat.gov.ru/', RosstatTLSAdapter())
                session.mount('https://www.rosstat.gov.ru/', RosstatTLSAdapter())
            response = session.get(url, params=params, timeout=30)
            response.raise_for_status()
            return response.content
    except requests.RequestException as error:
        logger.warning('Data source failed %s: %s', url, error)
        raise DataError(f'Источник {urlparse(url).hostname} недоступен. Повторите загрузку позже.') from error


def parse_cpi(content: bytes) -> pd.Series:
    """National all-items monthly CPI, never food/services or year-on-year CPI."""
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        sheet = workbook['01']
        rows = list(sheet.values)
        header = next(row for row in rows if sum(isinstance(v, (int, float)) and 1990 <= v <= 2100 for v in row) >= 5)
        years = {i: int(value) for i, value in enumerate(header) if isinstance(value, (int, float)) and 1990 <= value <= 2100}
        monthly = {}
        active = False
        for row in rows:
            label = str(row[0] or '').strip().lower()
            if label.startswith('к концу предыдущего месяца'):
                active = True
                continue
            if active and label.startswith('к декабрю'):
                break
            if active and label in MONTHS:
                month = MONTHS.index(label) + 1
                for col, year in years.items():
                    if col < len(row) and row[col] is not None:
                        value = float(str(row[col]).replace(',', '.'))
                        if not np.isfinite(value) or value <= 0:
                            raise ValueError('Invalid CPI')
                        monthly[pd.Period(f'{year}-{month:02}', freq='M').end_time.normalize()] = value / 100 - 1
        workbook.close()
        if not monthly:
            raise ValueError('Empty CPI')
        return pd.Series(monthly).sort_index()
    except (ValueError, KeyError, TypeError, StopIteration, OSError, BadZipFile) as error:
        raise DataError('Не удалось прочитать месячный ИПЦ Росстата. Формат источника изменился.') from error


@lru_cache(maxsize=2)
def _cpi_cached(bucket: int):
    soup = BeautifulSoup(_get(ROSSTAT_URL, rosstat=True), 'html.parser')
    item = next((node for node in soup.select('.document-list__item')
                 if 'Российской Федерации, месяцы' in node.get_text(' ', strip=True)), None)
    link = item.find('a', href=True) if item else None
    if not link:
        raise DataError('Росстат не предоставил ссылку на месячный ИПЦ Российской Федерации.')
    url = urljoin(ROSSTAT_URL, link['href'])
    if urlparse(url).hostname not in {'rosstat.gov.ru', 'www.rosstat.gov.ru'} or not urlparse(url).path.endswith('.xlsx'):
        raise DataError('Неожиданный адрес файла ИПЦ Росстата.')
    return parse_cpi(_get(url, rosstat=True)), url


def load_cpi():
    series, url = _cpi_cached(int(time.time() // 86400))
    return series.copy(), url


def parse_m2(content: bytes) -> pd.Series:
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        rows = list(workbook['Денежные агрегаты'].values)
        dates = rows[0]
        row = next(row for row in rows if re.fullmatch(r'Денежный агрегат [МM]2', str(row[0] or '').strip()))
        values = {}
        for period, value in zip(dates[1:], row[1:]):
            if isinstance(period, datetime) and isinstance(value, (int, float)) and np.isfinite(value) and value > 0:
                # Уровень на 1-е число отражает конец предыдущего месяца.
                values[pd.Timestamp(period).normalize() - pd.Timedelta(days=1)] = float(value)
        workbook.close()
        if not values:
            raise ValueError('Empty M2')
        return pd.Series(values).sort_index()
    except (ValueError, KeyError, TypeError, StopIteration, OSError, BadZipFile) as error:
        raise DataError('Не удалось прочитать денежный агрегат M2 Банка России.') from error


@lru_cache(maxsize=2)
def _m2_cached(bucket: int):
    return parse_m2(_get(M2_URL))


def load_m2():
    return _m2_cached(int(time.time() // 86400)).copy()


def parse_fx(content: bytes) -> pd.Series:
    try:
        root = ElementTree.fromstring(content)
        values = {}
        for record in root.findall('Record'):
            date = pd.to_datetime(record.attrib['Date'], format='%d.%m.%Y')
            nominal = float(record.findtext('Nominal').replace(',', '.'))
            quote = float(record.findtext('Value').replace(',', '.'))
            if nominal <= 0 or quote <= 0 or not np.isfinite(quote / nominal):
                raise ValueError('Invalid FX quote')
            values[date] = quote / nominal
        if not values:
            raise ValueError('Empty FX')
        return pd.Series(values).sort_index()
    except (ElementTree.ParseError, ValueError, KeyError, AttributeError, TypeError) as error:
        raise DataError('Банк России не вернул корректную историю курса валюты.') from error


@lru_cache(maxsize=64)
def _fx_cached(currency: str, start: str, end: str, bucket: int):
    params = {'date_req1': pd.Timestamp(start).strftime('%d/%m/%Y'),
              'date_req2': pd.Timestamp(end).strftime('%d/%m/%Y'), 'VAL_NM_RQ': FX_IDS[currency]}
    return parse_fx(_get(FX_URL, params=params))


def load_fx(currency, start, end):
    return _fx_cached(currency, start, end, int(time.time() // 86400)).copy()


@lru_cache(maxsize=128)
def _prices_cached(index_id: str, start: str, end: str, bucket: int):
    item = INDEX_BY_ID[index_id]
    if item['source'] == 'MOEX':
        series = _load_moex_prices(item['ticker'], start, end)
    else:
        exclusive_end = (pd.Timestamp(end) + pd.Timedelta(days=1)).date().isoformat()
        series = _load_yahoo_prices(item['ticker'], start, exclusive_end)
    series = series.copy()
    series.index = pd.to_datetime(series.index).tz_localize(None).normalize()
    return series[~series.index.duplicated(keep='last')].sort_index()


def load_index(index_id, start, end):
    return _prices_cached(index_id, start, end, int(time.time() // 86400)).copy()


def convert_monthly_to_rubles(prices: pd.Series, fx: pd.Series | None) -> pd.Series:
    monthly = prices.resample('ME').last()
    if fx is not None:
        # Стоимость на конец календарного месяца по последнему действующему курсу.
        rates = fx.sort_index().reindex(monthly.index, method='ffill')
        monthly = monthly * rates
    return monthly


def load_history(indices, start, end, include_m2=True):
    today = pd.Timestamp(datetime.now(ZoneInfo('Europe/Moscow')).date())
    start_month = pd.Timestamp(start).to_period('M')
    last_complete = today.to_period('M') - 1
    requested_end = pd.Timestamp(end or today).to_period('M')
    last_month = min(requested_end, last_complete)
    if last_month < start_month:
        raise DataError('Выберите хотя бы два завершённых месяца.')
    baseline = (start_month - 1).end_time.normalize()
    final = last_month.end_time.normalize()
    fetch_start = (baseline - pd.Timedelta(days=40)).date().isoformat()
    fetch_end = final.date().isoformat()
    warnings = []
    cpi, cpi_source = load_cpi()
    last_with_cpi = cpi.index.max().to_period('M')
    if last_month > last_with_cpi:
        last_month = last_with_cpi
        final = last_month.end_time.normalize()
        warnings.append(f'Период ограничен {final.date()}: более свежий ИПЦ ещё не опубликован.')
    if final <= baseline:
        raise DataError('За выбранный период ещё нет опубликованного ИПЦ.')
    fx = {}
    for currency in sorted({INDEX_BY_ID[key]['currency'] for key in indices} - {'RUB'}):
        fx[currency] = load_fx(currency, fetch_start, fetch_end)
    from concurrent.futures import ThreadPoolExecutor
    moex_ids = [key for key in indices if INDEX_BY_ID[key]['source'] == 'MOEX']
    levels = {}
    if moex_ids:
        with ThreadPoolExecutor(max_workers=min(4, len(moex_ids))) as pool:
            levels.update(zip(moex_ids, pool.map(lambda key: load_index(key, fetch_start, fetch_end), moex_ids)))
    for key in indices:
        if key not in levels:
            levels[key] = load_index(key, fetch_start, fetch_end)
    prices = pd.DataFrame({key: convert_monthly_to_rubles(levels[key], fx.get(INDEX_BY_ID[key]['currency'])) for key in indices})
    dates = pd.date_range(baseline, final, freq='ME')
    prices = prices.reindex(dates)
    valid = prices.notna().all(axis=1)
    if not valid.any():
        raise DataError('У выбранных индексов нет общей истории за этот период.')
    first, last = np.flatnonzero(valid)[[0, -1]]
    if first or last != len(prices) - 1:
        warnings.append('Период сокращён до общей доступной истории выбранных индексов.')
    prices = prices.iloc[first:last + 1]
    if prices.isna().any().any():
        raise DataError('В общей истории есть пропущенный месяц. Измените период или состав индексов.')
    returns = prices.pct_change(fill_method=None).iloc[1:]
    inflation = cpi.reindex(returns.index)
    if inflation.isna().any():
        raise DataError('Не хватает месячных значений ИПЦ за выбранный период.')
    if len(returns) < 2:
        raise DataError('Для расчёта риска нужно минимум два месяца общей истории.')
    if len(returns) < 30:
        warnings.append('История короче 30 месяцев: оценки доходности и риска менее устойчивы.')
    money = None
    if include_m2:
        try:
            money = load_m2().reindex(prices.index)
            if pd.isna(money.iloc[0]):
                warnings.append('Для начала периода нет значения M2; линия сравнения недоступна.')
                money = None
            elif money.isna().any():
                warnings.append('M2 показан только до последнего опубликованного месяца; пропуски не заполняются.')
        except DataError as error:
            warnings.append(str(error) + ' Портфель рассчитан без линии M2.')
    if any(INDEX_BY_ID[key]['return_type'] == 'price' for key in indices):
        warnings.append('Ценовые индексы не учитывают дивиденды. Индексы полной доходности учитывают выплаты по своей методике.')
    if fx:
        warnings.append('Зарубежные индексы пересчитаны в рубли по официальным историческим курсам Банка России.')
    return {'returns': returns, 'inflation': inflation, 'm2': money, 'baseline': prices.index[0],
            'warnings': warnings, 'cpi_source': cpi_source}
