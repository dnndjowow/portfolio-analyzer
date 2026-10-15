from datetime import datetime
import io

import numpy as np
import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from data_loader import DataError
from index_analysis import analyze_history, simulate_portfolio
from index_catalog import INDEX_CATALOG
from main import app
from market_indices import convert_monthly_to_rubles, parse_cpi, parse_fx, parse_m2
from schemas.indices import IndexPortfolioRequest


def history(n=36):
    dates = pd.date_range('2020-01-31', periods=n, freq='ME')
    values = np.column_stack([0.01 + 0.04 * np.sin(np.arange(n)), 0.005 + 0.015 * np.cos(np.arange(n))])
    returns = pd.DataFrame(values, columns=['MCFTR', 'RGBITR'], index=dates)
    return {'returns': returns, 'inflation': pd.Series(0.005, index=dates),
            'm2': pd.Series(np.linspace(100, 200, n + 1), index=pd.DatetimeIndex([pd.Timestamp('2019-12-31'), *dates])),
            'baseline': pd.Timestamp('2019-12-31'), 'warnings': [], 'cpi_source': 'https://rosstat.gov.ru/test.xlsx'}


def request(**kwargs):
    return IndexPortfolioRequest(indices=['MCFTR', 'RGBITR'], **kwargs)


def test_buy_and_hold_differs_from_constant_weight_rebalancing():
    returns = pd.DataFrame([[1, 0], [0, 1]], index=pd.to_datetime(['2020-02-29', '2020-03-31']))
    weights = np.array([0.5, 0.5])
    _, hold, start_weights, _ = simulate_portfolio(returns, weights, 'none')
    _, monthly, _, _ = simulate_portfolio(returns, weights, 'monthly')
    assert hold.tolist() == pytest.approx([1, 1.5, 2])
    assert monthly.tolist() == pytest.approx([1, 1.5, 2.25])
    assert start_weights[1] == pytest.approx([2 / 3, 1 / 3])


@pytest.mark.parametrize('dates,mode,expected', [
    (['2020-02-29', '2020-03-31'], 'quarterly', 2),
    (['2020-03-31', '2020-04-30'], 'quarterly', 2.25),
    (['2020-03-31', '2020-04-30'], 'yearly', 2),
    (['2020-12-31', '2021-01-31'], 'yearly', 2.25),
])
def test_rebalance_uses_calendar_boundaries(dates, mode, expected):
    returns = pd.DataFrame([[1, 0], [0, 1]], index=pd.to_datetime(dates))
    _, path, _, _ = simulate_portfolio(returns, np.array([0.5, 0.5]), mode)
    assert path[-1] == pytest.approx(expected)


@pytest.mark.parametrize('rebalance', ['none', 'monthly', 'quarterly', 'yearly'])
def test_real_curve_is_nominal_curve_divided_by_cumulative_cpi(rebalance):
    h = history()
    result = analyze_history(h, request(rebalance=rebalance))
    assert result['final_real_capital'] == pytest.approx(result['final_nominal_capital'] / np.prod(1 + h['inflation']))
    assert result['rows'][0]['nominal'] == pytest.approx(100000)
    assert result['rows'][0]['real'] == pytest.approx(100000)
    assert result['rows'][0]['m2_growth'] == pytest.approx(100)


@pytest.mark.parametrize('frequency', ['monthly', 'quarterly', 'yearly'])
def test_graph_frequency_does_not_change_strategy_or_monthly_risk(frequency):
    monthly = analyze_history(history(), request(rebalance='quarterly'))
    other = analyze_history(history(), request(rebalance='quarterly', frequency=frequency))
    assert other['nominal_metrics'] == monthly['nominal_metrics']
    assert other['real_metrics'] == monthly['real_metrics']
    assert other['final_weights'] == monthly['final_weights']
    assert other['rows'][-1]['real'] == monthly['rows'][-1]['real']
    assert len(other['rows']) == {'monthly': 37, 'quarterly': 13, 'yearly': 4}[frequency]
    period_nominal = np.prod([1 + row['nominal_return'] for row in other['rows'][1:]]) - 1
    assert period_nominal == pytest.approx(other['nominal_metrics']['total_return'])
    for row in other['rows'][1:]:
        assert row['real_return'] == pytest.approx((1 + row['nominal_return']) / (1 + row['inflation']) - 1)


def test_annual_graph_keeps_intrayear_drawdown_and_volatility_in_metrics():
    h = history(12)
    h['returns'].iloc[:] = 0
    h['returns'].iloc[0] = -0.5
    h['returns'].iloc[1] = 1
    result = analyze_history(h, request(frequency='yearly'))
    assert result['nominal_metrics']['max_drawdown'] == pytest.approx(-0.5)
    assert result['nominal_metrics']['volatility'] > 0
    assert result['nominal_metrics']['total_return'] == pytest.approx(0)


def test_capital_only_scales_money_not_return_or_risk():
    a = analyze_history(history(), request(capital=100000))
    b = analyze_history(history(), request(capital=200000))
    assert a['weights'] == b['weights']
    assert a['nominal_metrics'] == b['nominal_metrics']
    assert a['real_metrics'] == b['real_metrics']
    assert b['final_nominal_capital'] == pytest.approx(a['final_nominal_capital'] * 2)
    assert a['rows'][-1]['nominal_growth'] == b['rows'][-1]['nominal_growth']


@pytest.mark.parametrize('mode', ['manual', 'min_vol', 'max_sharpe'])
def test_m2_never_changes_weights_returns_risk_or_optimization(mode):
    h = history()
    a = analyze_history(h, request(mode=mode))
    h['m2'] = h['m2'] * np.linspace(1, 1000, len(h['m2']))
    b = analyze_history(h, request(mode=mode))
    assert a['weights'] == b['weights']
    assert a['nominal_metrics'] == b['nominal_metrics']
    assert a['real_metrics'] == b['real_metrics']
    assert a['rows'][-1]['m2_growth'] != b['rows'][-1]['m2_growth']
    assert len(b['weights']) == 2
    assert min(b['weights']) >= 0


def test_foreign_index_return_in_rubles_compounds_fx_return():
    prices = pd.Series([100, 110], index=pd.to_datetime(['2020-01-30', '2020-02-28']))
    fx = pd.Series([60, 66], index=pd.to_datetime(['2020-01-31', '2020-02-29']))
    rubles = convert_monthly_to_rubles(prices, fx)
    assert rubles.iloc[-1] / rubles.iloc[0] - 1 == pytest.approx(1.1 * 1.1 - 1)


def test_fx_uses_nominal_units_and_effective_date():
    xml = b'<ValCurs><Record Date="31.01.2020"><Nominal>100</Nominal><Value>60,00</Value></Record><Record Date="28.02.2020"><Nominal>100</Nominal><Value>70,00</Value></Record></ValCurs>'
    series = parse_fx(xml)
    assert series.tolist() == [0.6, 0.7]
    assert series.index[0] == pd.Timestamp('2020-01-31')


def workbook_bytes(workbook):
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_m2_parser_uses_national_m2_not_m2x_and_shifts_first_day():
    w = Workbook(); s = w.active; s.title = 'Денежные агрегаты'
    s.append(['Денежные агрегаты', datetime(2020, 1, 1), datetime(2020, 2, 1)])
    s.append(['Денежный агрегат M2X', 999, 9999])
    s.append(['Денежный агрегат М2', 100, 120])
    m2 = parse_m2(workbook_bytes(w))
    assert m2.tolist() == [100, 120]
    assert m2.index.tolist() == list(pd.to_datetime(['2019-12-31', '2020-01-31']))


def test_cpi_parser_uses_monthly_national_all_items_sheet():
    w = Workbook(); s = w.active; s.title = '01'
    s.append(['Индексы потребительских цен по Российской Федерации'])
    s.append([None, 2019, 2020, 2021, 2022, 2023])
    s.append(['к концу предыдущего месяца'])
    s.append(['январь', 101, 102, 103, 104, 105])
    s.append(['февраль', 99, 100, 101, 102, 103])
    s.append(['к декабрю предыдущего года'])
    s.append(['декабрь', 999, 999, 999, 999, 999])
    other = w.create_sheet('02'); other.append(['январь', 150])
    cpi = parse_cpi(workbook_bytes(w))
    assert cpi.loc['2020-01-31'] == pytest.approx(0.02)
    assert cpi.loc['2019-02-28'] == pytest.approx(-0.01)
    assert len(cpi) == 10


client = TestClient(app)


def test_catalog_has_exactly_ten_indices_and_no_money_supply():
    response = client.get('/api/index-catalog')
    assert response.status_code == 200
    items = response.json()['indices']
    assert len(items) == 10
    assert not any('M2' in item['id'] for item in items)
    assert len({item['id'] for item in items}) == 10


@pytest.mark.parametrize('extra', [
    {'indices': ['MCFTR']}, {'indices': ['MCFTR', 'M2']}, {'indices': ['MCFTR', 'MCFTR']},
    {'weights': [-0.1, 1.1]}, {'weights': [0.4, 0.4]}, {'weights': [1]},
    {'allow_short': True}, {'capital': 0}, {'capital': -100}, {'base_currency': 'USD'},
    {'start': '2020-01-15'}, {'start': '2020-01-01', 'end': '2019-12-31'},
])
def test_invalid_requests_are_rejected_before_sources_are_called(extra, monkeypatch):
    import routers.indices as router
    monkeypatch.setattr(router, 'load_history', lambda *args: pytest.fail('Unexpected data request'))
    response = client.post('/api/index-portfolio', json={'indices': ['MCFTR', 'RGBITR'], **extra})
    assert response.status_code == 422


def test_new_endpoint_returns_both_curves_and_separate_m2(monkeypatch):
    import routers.indices as router
    monkeypatch.setattr(router, 'load_history', lambda *args: history())
    response = client.post('/api/index-portfolio', json={'indices': ['MCFTR', 'RGBITR'], 'weights': [0.25, 0.75], 'capital': 100000, 'rebalance': 'none', 'frequency': 'quarterly'})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['weights'] == [0.25, 0.75]
    assert len(result['indices']) == 2
    assert len(result['rows']) == 13
    assert result['rows'][0]['m2_growth'] == 100
    assert 'nominal_metrics' in result and 'real_metrics' in result


def test_source_failure_is_explicit_without_constant_inflation_fallback(monkeypatch):
    import routers.indices as router
    def fail(*args):
        raise DataError('ИПЦ Росстата недоступен')
    monkeypatch.setattr(router, 'load_history', fail)
    response = client.post('/api/index-portfolio', json={'indices': ['MCFTR', 'RGBITR']})
    assert response.status_code == 422
    assert 'ИПЦ Росстата недоступен' in response.json()['detail']


def test_automatic_weights_require_sufficient_history_but_manual_analysis_is_available():
    h = history(12)
    assert analyze_history(h, request())['n_months'] == 12
    with pytest.raises(DataError, match='минимум 30'):
        analyze_history(h, request(mode='min_vol'))


def mock_market(monkeypatch):
    import market_indices as market
    dates = pd.date_range('2019-12-31', periods=5, freq='ME')
    monkeypatch.setattr(market, 'load_cpi', lambda: (pd.Series(0.01, index=dates[1:]), 'https://rosstat.gov.ru/test.xlsx'))
    monkeypatch.setattr(market, 'load_index', lambda *args: pd.Series([100, 110, 105, 112, 115], index=dates))
    return market, dates


def test_unpublished_m2_months_do_not_shorten_portfolio_history(monkeypatch):
    market, dates = mock_market(monkeypatch)
    monkeypatch.setattr(market, 'load_m2', lambda: pd.Series([100, 110, 120], index=dates[:3]))
    h = market.load_history(['MCFTR', 'RGBITR'], '2020-01-01', '2020-04-30')
    assert len(h['returns']) == 4
    assert pd.isna(h['m2'].iloc[-1])
    assert any('не заполняются' in warning for warning in h['warnings'])


def test_m2_can_be_disabled_without_loading_it(monkeypatch):
    market, _ = mock_market(monkeypatch)
    monkeypatch.setattr(market, 'load_m2', lambda: pytest.fail('M2 should not be requested'))
    h = market.load_history(['MCFTR', 'RGBITR'], '2020-01-01', '2020-04-30', include_m2=False)
    assert h['m2'] is None
    assert len(h['returns']) == 4


def test_missing_internal_month_is_not_filled(monkeypatch):
    market, dates = mock_market(monkeypatch)
    monkeypatch.setattr(market, 'load_index', lambda *args: pd.Series([100, 110, 112, 115], index=dates.delete(2)))
    with pytest.raises(DataError, match='пропущенный месяц'):
        market.load_history(['MCFTR', 'RGBITR'], '2020-01-01', '2020-04-30', include_m2=False)


def test_catalog_uses_current_corporate_bond_index():
    assert any(item['id'] == 'RUCBTRNS' for item in INDEX_CATALOG)
    assert not any(item['id'] == 'RUCBITR' for item in INDEX_CATALOG)


@pytest.mark.parametrize('parser', [parse_cpi, parse_m2])
def test_invalid_macro_workbook_produces_a_readable_source_error(parser):
    with pytest.raises(DataError):
        parser(b'<html>maintenance</html>')
