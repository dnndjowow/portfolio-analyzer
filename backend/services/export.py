import io

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font

import metrics
from schemas.portfolio import ExportRequest
from services.portfolio import prepare_portfolio, check_weights, get_curve_index


def create_portfolio_export(data: ExportRequest) -> io.BytesIO:
    returns_frame, portfolio_inputs, periods_per_year, warnings = prepare_portfolio(data)
    workbook = Workbook()

    weights_sheet = workbook.active
    weights_sheet.title = 'Веса'
    weights_sheet.append(['Портфель', *returns_frame.columns])

    for portfolio in data.portfolios:
        portfolio_weights = check_weights(portfolio.weights, data.allow_short)
        weights_sheet.append([portfolio.name, *portfolio_weights.tolist()])

    metrics_sheet = workbook.create_sheet('Показатели')
    metrics_sheet.append([
        'Портфель',
        'Ожид. доход (год)',
        'CAGR',
        'Волатильность (год)',
        'Шарп',
        'Макс. просадка',
        'Макс. период восст., периодов',
        'Восстановился',
    ])

    curves_sheet = workbook.create_sheet('Кривые')
    portfolio_names = [portfolio.name for portfolio in data.portfolios]
    curves_sheet.append(['Период', *portfolio_names])
    portfolio_curves = []

    for portfolio in data.portfolios:
        portfolio_metrics = metrics.compute_metrics(
            returns_frame.values,
            np.asarray(portfolio.weights),
            periods_per_year,
            data.risk_free,
        )

        recovery_status = 'нет'
        if portfolio_metrics['recovered']:
            recovery_status = 'да'

        metrics_sheet.append([
            portfolio.name,
            portfolio_metrics['expected_return'],
            portfolio_metrics['cagr'],
            portfolio_metrics['volatility'],
            portfolio_metrics['sharpe'],
            portfolio_metrics['max_drawdown'],
            portfolio_metrics['max_recovery_periods'],
            recovery_status,
        ])
        portfolio_curves.append(portfolio_metrics['curve'])

    curve_index = get_curve_index(returns_frame)

    for period_index, period_label in enumerate(curve_index):
        curve_values = [curve[period_index] for curve in portfolio_curves]
        curves_sheet.append([period_label, *curve_values])

    returns_sheet = workbook.create_sheet('Реальные доходности')
    returns_sheet.append(['Период', *returns_frame.columns])

    for period_label, row in zip(curve_index[1:], returns_frame.values.tolist()):
        returns_sheet.append([period_label, *row])

    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)

    for row in weights_sheet.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = '0.00%'

    for row in metrics_sheet.iter_rows(min_row=2, min_col=2, max_col=6):
        for cell in row:
            if cell.column == 5:
                cell.number_format = '0.00'
            else:
                cell.number_format = '0.00%'

    export_file = io.BytesIO()
    workbook.save(export_file)
    export_file.seek(0)

    return export_file
