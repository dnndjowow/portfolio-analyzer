"""FastAPI-приложение «Портфель»."""
from __future__ import annotations

import io
import logging
import os
import time
from pathlib import Path
from typing import Literal

import numpy as np
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from dotenv import load_dotenv

import data_loader as dl
import metrics as mt
import optimizer as opt
import assistant_service as assistant

load_dotenv(Path(__file__).resolve().parents[1] / ".env")

logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO"),
                    format="%(asctime)s %(levelname)s %(name)s: %(message)s")
log = logging.getLogger("portfolio.api")

app = FastAPI(title="Портфель API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[o.strip() for o in os.getenv("CORS_ORIGINS", "*").split(",")],
    allow_methods=["*"], allow_headers=["*"],
)


# ---------------------------------------------------------------------------
class Base(BaseModel):
    returns: dl.Dataset
    settings: dl.Settings = dl.Settings()
    risk_free: float = Field(0.05, description="Годовая безрисковая ставка в долях")
    allow_short: bool = False


class OptimizeRequest(Base):
    mode: opt.Mode
    target_return: float | None = None
    target_vol: float | None = None


class FrontierRequest(Base):
    n_points: int = Field(40, ge=5, le=200)
    n_samples: int = Field(5000, ge=100, le=20000)


class MetricsRequest(Base):
    weights: list[float]


class NamedWeights(BaseModel):
    name: str
    weights: list[float]


class CompareRequest(Base):
    portfolios: list[NamedWeights] = Field(..., min_length=1, max_length=5)


class ExportRequest(Base):
    portfolios: list[NamedWeights] = Field(..., min_length=1)


class TickersRequest(BaseModel):
    tickers: list[str]
    start: str = "2015-01-01"
    end: str | None = None
    frequency: str = "monthly"
    m2_levels: list[dict] | None = None


class SupportRequest(BaseModel):
    stage: Literal["upload", "tickers", "frontier", "compare", "optimize", "export"]
    error: str = Field(..., min_length=1, max_length=800)
    context: dict = Field(default_factory=dict)


_assistant_calls: dict[str, list[float]] = {}


def _prepare(req: Base):
    if not np.isfinite(req.risk_free):
        raise HTTPException(422, "Безрисковая ставка должна быть конечным числом.")
    try:
        df, warnings = dl.preprocess(req.returns, req.settings)
    except dl.DataError as e:
        raise HTTPException(422, str(e)) from e
    ppy = req.settings.periods_per_year
    inp = opt.make_inputs(df.values, ppy, req.allow_short)
    rf = dl.real_rate(req.risk_free, req.settings) if req.settings.risk_free_nominal else req.risk_free
    req.risk_free = rf  # дальше везде реальная ставка
    return df, inp, ppy, warnings


def _check_weights(w: list[float], allow_short: bool = False) -> np.ndarray:
    arr = np.asarray(w, dtype=float)
    if arr.shape != (dl.N_ASSETS,):
        raise HTTPException(422, f"Нужно {dl.N_ASSETS} весов, получено {arr.size}.")
    if not np.isfinite(arr).all():
        raise HTTPException(422, "Веса должны быть конечными числами.")
    total = float(arr.sum())
    if abs(total - 1) > 1e-6:
        raise HTTPException(422, f"Сумма весов должна быть 1, получено {total:.4f}.")
    if allow_short:
        if np.any(arr < -opt.SHORT_BOUND - 1e-8) or np.any(arr > 1 + 1e-8):
            raise HTTPException(422, f"При коротких позициях вес каждого актива должен быть от {-opt.SHORT_BOUND:.0%} до 100%.")
    elif np.any(arr < -1e-8):
        raise HTTPException(422, "Отрицательные веса допустимы только при включённых коротких позициях.")
    return arr


def _curve_index(df) -> list[str]:
    idx = [str(i.date()) if hasattr(i, "date") else str(i) for i in df.index]
    return ["старт", *idx]


# ---------------------------------------------------------------------------
@app.get("/api/health")
def health():
    return {"status": "ok"}


@app.post("/api/assistant/diagnose", include_in_schema=False)
@app.post("/api/support/diagnose")
def support_diagnose(req: SupportRequest, request: Request):
    if not os.getenv("OPENROUTER_API_KEY"):
        log.warning("Support diagnosis requested without OPENROUTER_API_KEY")
        raise HTTPException(503, "Автоматическая диагностика временно недоступна.")
    client = request.client.host if request.client else "unknown"
    now = time.monotonic()
    recent = [t for t in _assistant_calls.get(client, []) if now - t < 3600]
    if recent and now - recent[-1] < 15:
        raise HTTPException(429, "Проверка недавней ошибки ещё выполняется. Попробуйте чуть позже.")
    if len(recent) >= 20:
        raise HTTPException(429, "Слишком много запросов диагностики. Попробуйте позже.")
    recent.append(now)
    _assistant_calls[client] = recent
    try:
        return {"analysis": assistant.diagnose(req.stage, req.error, req.context)}
    except assistant.AssistantError as e:
        log.warning("Support diagnosis unavailable: %s", e)
        raise HTTPException(502, "Автоматическая диагностика временно недоступна.") from e


@app.post("/api/upload")
async def upload(file: UploadFile = File(...)):
    content = await file.read()
    name = (file.filename or "").lower()
    try:
        ds, warnings = dl.parse_csv(content) if name.endswith(".csv") else dl.parse_excel(content)
        ds, w2 = dl.ensure_money_supply(ds)
        warnings += w2
        if len(ds.columns) != dl.N_ASSETS:
            raise dl.DataError("Ожидается 10 индикаторов.")
        if len(ds.data) < dl.MIN_OBS:
            warnings.append(f"В файле {len(ds.data)} наблюдений — нужно не меньше {dl.MIN_OBS}.")
    except dl.DataError as e:
        raise HTTPException(422, str(e)) from e
    except Exception as e:  # битый файл
        log.exception("upload failed")
        raise HTTPException(422, f"Не удалось прочитать файл: {e}") from e
    log.info("upload %s: %d rows × %d cols", file.filename, len(ds.data), len(ds.columns))
    return {"dataset": ds, "preview": {"index": ds.index[:20], "data": ds.data[:20]},
            "n_obs": len(ds.data), "suggested_units": dl.guess_units(ds), "warnings": warnings}


@app.post("/api/load-tickers")
def load_tickers(req: TickersRequest):
    try:
        ds, warnings = dl.load_tickers(req.tickers, req.start, req.end, req.frequency, req.m2_levels)
    except dl.DataError as e:
        raise HTTPException(422, str(e)) from e
    return {"dataset": ds, "preview": {"index": ds.index[:20], "data": ds.data[:20]},
            "n_obs": len(ds.data), "suggested_units": "percent", "warnings": warnings}


@app.post("/api/optimize")
def optimize(req: OptimizeRequest):
    df, inp, ppy, warnings = _prepare(req)
    try:
        w = opt.optimize(inp, req.mode, req.risk_free, req.target_return, req.target_vol)
    except opt.OptimizationError as e:
        raise HTTPException(422, str(e)) from e
    if req.mode == opt.Mode.MAX_SHARPE and (msg := opt.sharpe_warning(inp, req.risk_free)):
        warnings.append(msg)
    m = mt.compute_metrics(df.values, w, ppy, req.risk_free)
    return {
        "mode": req.mode, "risk_free_real": req.risk_free, "assets": list(df.columns), "weights": w.tolist(),
        "point": {"ret": opt.port_return(w, inp), "vol": opt.port_vol(w, inp)},
        "metrics": m, "index": _curve_index(df), "warnings": warnings,
    }


@app.post("/api/efficient-frontier")
def frontier(req: FrontierRequest):
    df, inp, ppy, warnings = _prepare(req)
    try:
        w_mv = opt.min_volatility(inp)
        special = {"min_vol": {"ret": opt.port_return(w_mv, inp), "vol": opt.port_vol(w_mv, inp),
                               "weights": w_mv.tolist()}}
        w_ms = opt.max_sharpe(inp, req.risk_free)
        special["max_sharpe"] = {"ret": opt.port_return(w_ms, inp), "vol": opt.port_vol(w_ms, inp),
                                 "weights": w_ms.tolist(), "sharpe": opt.sharpe(w_ms, inp, req.risk_free)}
        if msg := opt.sharpe_warning(inp, req.risk_free):
            warnings.append(msg)
        pts = opt.efficient_frontier(inp, req.n_points)
    except opt.OptimizationError as e:
        raise HTTPException(422, str(e)) from e
    assets = [{"name": c, "ret": float(inp.mu[i]), "vol": float(np.sqrt(inp.cov[i, i]))}
              for i, c in enumerate(df.columns)]
    return {"frontier": pts, "cloud": opt.monte_carlo(inp, req.risk_free, req.n_samples),
            "assets": assets, **special, "risk_free_real": req.risk_free, "warnings": warnings}


@app.post("/api/portfolio/metrics")
def portfolio_metrics(req: MetricsRequest):
    df, inp, ppy, warnings = _prepare(req)
    w = _check_weights(req.weights, req.allow_short)
    return {"metrics": mt.compute_metrics(df.values, w, ppy, req.risk_free),
            "point": {"ret": opt.port_return(w, inp), "vol": opt.port_vol(w, inp)},
            "index": _curve_index(df), "warnings": warnings}


@app.post("/api/compare")
def compare(req: CompareRequest):
    df, inp, ppy, warnings = _prepare(req)
    series = []
    for p in req.portfolios:
        w = _check_weights(p.weights, req.allow_short)
        m = mt.compute_metrics(df.values, w, ppy, req.risk_free)
        series.append({"name": p.name, "curve": m.pop("curve"), "drawdown": m.pop("drawdown"),
                       "point": {"ret": opt.port_return(w, inp), "vol": opt.port_vol(w, inp)},
                       "metrics": {k: v for k, v in m.items() if k != "period_returns"}})
    return {"index": _curve_index(df), "series": series, "warnings": warnings}


@app.post("/api/export")
def export(req: ExportRequest):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    df, inp, ppy, _ = _prepare(req)
    wb = Workbook()
    ws = wb.active
    ws.title = "Веса"
    ws.append(["Портфель", *df.columns])
    for p in req.portfolios:
        ws.append([p.name, *_check_weights(p.weights, req.allow_short).tolist()])
    ws2 = wb.create_sheet("Показатели")
    heads = ["Портфель", "Ожид. доход (год)", "CAGR", "Волатильность (год)", "Шарп",
             "Макс. просадка", "Макс. период восст., периодов", "Восстановился"]
    ws2.append(heads)
    ws3 = wb.create_sheet("Кривые")
    ws3.append(["Период", *[p.name for p in req.portfolios]])
    curves = []
    for p in req.portfolios:
        m = mt.compute_metrics(df.values, np.asarray(p.weights), ppy, req.risk_free)
        ws2.append([p.name, m["expected_return"], m["cagr"], m["volatility"], m["sharpe"],
                    m["max_drawdown"], m["max_recovery_periods"], "да" if m["recovered"] else "нет"])
        curves.append(m["curve"])
    for i, label in enumerate(_curve_index(df)):
        ws3.append([label, *[c[i] for c in curves]])
    ws4 = wb.create_sheet("Реальные доходности")
    ws4.append(["Период", *df.columns])
    for label, row in zip(_curve_index(df)[1:], df.values.tolist()):
        ws4.append([label, *row])
    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=2):
        for c in row:
            c.number_format = "0.00%"
    for row in ws2.iter_rows(min_row=2, min_col=2, max_col=6):
        for c in row:
            c.number_format = "0.00%" if c.column != 5 else "0.00"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="portfolio_export.xlsx"'})
