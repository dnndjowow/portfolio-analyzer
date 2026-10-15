import io
import logging
import re
from datetime import date, datetime
from concurrent.futures import ThreadPoolExecutor

import numpy as np
import pandas as pd

from schemas.data import Dataset, InflationSettings, Settings
from config import (
    MIN_ASSETS,
    MAX_ASSETS,
    MIN_OBS,
    PERIODS_PER_YEAR,
    RESAMPLE_RULE,
    EXCLUDED_INDICATOR_RE,
)


logger = logging.getLogger('portfolio.data')


class DataError(ValueError):
    pass


def _is_number(value) -> bool:
    if not isinstance(value, (int, float, np.number)) or isinstance(value, bool):
        return False

    return np.isfinite(value)


def parse_excel(content: bytes) -> tuple[Dataset, list[str]]:
    from openpyxl import load_workbook

    values_sheet = load_workbook(
        io.BytesIO(content),
        data_only=True,
        read_only=True,
    ).worksheets[0]
    formulas_sheet = load_workbook(
        io.BytesIO(content),
        data_only=False,
        read_only=True,
    ).worksheets[0]
    rows = []

    for values_row, formulas_row in zip(
        values_sheet.iter_rows(values_only=True),
        formulas_sheet.iter_rows(values_only=True),
    ):
        row = []

        for value, formula in zip(values_row, formulas_row):
            # Итоговые строки с формулами не являются наблюдениями.
            if isinstance(formula, str) and formula.startswith('='):
                row.append(None)
            else:
                row.append(value)

        rows.append(row)

    return _parse_grid(rows)


def parse_csv(content: bytes) -> tuple[Dataset, list[str]]:
    text = content.decode('utf-8-sig', errors='replace')
    separator = ','
    decimal_separator = '.'

    if text.count(';') > text.count(','):
        separator = ';'
        decimal_separator = ','

    data_frame = pd.read_csv(
        io.StringIO(text),
        sep=separator,
        header=None,
        decimal=decimal_separator,
    )
    raw_rows = data_frame.where(pd.notna(data_frame), None).values.tolist()
    rows = []

    for row in raw_rows:
        rows.append([_parse_value(value) for value in row])

    return _parse_grid(rows)


def _parse_value(value):
    if isinstance(value, str):
        try:
            return float(value.replace(',', '.').replace(' ', ''))
        except ValueError:
            return value

    return value


def _parse_grid(rows: list[list]) -> tuple[Dataset, list[str]]:
    header_row = None
    header_column = None
    columns = []
    period_headers = {'date', 'dates', 'дата', 'даты', 'период', 'period'}

    for row_index, row in enumerate(rows):
        first_column = 1 if row and str(row[0]).strip().lower() in period_headers else 0
        for column_index in range(first_column, len(row)):
            names = []
            for value in row[column_index:]:
                if not isinstance(value, str) or not value.strip():
                    break
                if value.strip().lower() in {'портфель', 'portfolio'}:
                    break
                names.append(value.strip())
            if len(names) >= MIN_ASSETS:
                header_row, header_column, columns = row_index, column_index, names
                break
        if header_row is not None:
            break

    if header_row is None:
        raise DataError('Не найдена строка заголовков с названиями активов (минимум 2).')
    n_assets = len(columns)
    if n_assets > MAX_ASSETS:
        raise DataError(f'Можно загрузить не более {MAX_ASSETS} активов.')
    if any(EXCLUDED_INDICATOR_RE.search(name) for name in columns):
        raise DataError('Показатель M2 исключён. Удалите его столбец из файла.')

    periods = []
    observations = []
    has_dates = False

    for row in rows[header_row + 1:]:
        padded_row = row + [None] * (header_column + n_assets)
        row_values = padded_row[header_column:header_column + n_assets]
        numeric_values = []

        for value in row_values:
            if _is_number(value):
                numeric_values.append(value)
            else:
                numeric_values.append(None)

        if all(value is None for value in numeric_values):
            if observations:
                break
            continue

        if sum(value is None for value in numeric_values) > n_assets // 2:
            break

        period_label = None
        if header_column > 0:
            period_label = row[header_column - 1]

        if isinstance(period_label, (datetime, date)):
            has_dates = True
            periods.append(pd.Timestamp(period_label).date().isoformat())
        elif isinstance(period_label, str) and _looks_like_date(period_label):
            has_dates = True
            periods.append(pd.Timestamp(period_label).date().isoformat())
        elif _is_number(period_label):
            periods.append(str(int(period_label)))
        else:
            periods.append(str(len(observations) + 1))

        observation = []
        for value in numeric_values:
            if value is None:
                observation.append(None)
            else:
                observation.append(float(value))

        observations.append(observation)

    dataset = Dataset(
        columns=columns,
        index=periods,
        data=observations,
        has_dates=has_dates,
    )

    return dataset, []


def _looks_like_date(value: str) -> bool:
    try:
        pd.Timestamp(value)
        return bool(re.search(r'\d{4}|\d{1,2}[./-]\d{1,2}', value))
    except (ValueError, TypeError):
        return False


def guess_units(dataset: Dataset) -> str:
    numeric_values = [
        value
        for row in dataset.data
        for value in row
        if value is not None
    ]
    returns_array = np.array(numeric_values, dtype=float)

    if returns_array.size and np.nanmax(np.abs(returns_array)) > 1.0:
        return 'percent'

    return 'fraction'


def to_frame(dataset: Dataset) -> pd.DataFrame:
    returns_frame = pd.DataFrame(dataset.data, columns=dataset.columns, dtype=float)

    if dataset.has_dates:
        returns_frame.index = pd.to_datetime(dataset.index)
    else:
        returns_frame.index = pd.Index(dataset.index, name='period')

    return returns_frame


def preprocess(dataset: Dataset, settings: Settings) -> tuple[pd.DataFrame, list[str]]:
    warnings = []
    returns_frame = to_frame(dataset)
    unit_divisor = 1.0

    if settings.units == 'percent':
        unit_divisor = 100.0

    returns_frame = returns_frame / unit_divisor
    missing_count = int(returns_frame.isna().sum().sum())

    if missing_count:
        if settings.missing == 'ffill':
            returns_frame = returns_frame.ffill().dropna()
            missing_method = 'ffill'
        else:
            returns_frame = returns_frame.dropna()
            missing_method = 'удаление строк'

        warnings.append(f'Обработано пропусков: {missing_count} ({missing_method}).')

    source_periods_per_year = PERIODS_PER_YEAR[settings.frequency]

    if settings.inflation.source == 'series':
        inflation_values = settings.inflation.series or []
        if len(inflation_values) != len(dataset.index):
            raise DataError(
                f'Ряд инфляции: {len(inflation_values)} значений, '
                f'а наблюдений {len(dataset.index)}.'
            )

        inflation_series = pd.Series(
            np.asarray(inflation_values, dtype=float) / unit_divisor,
            index=to_frame(dataset).index,
        ).reindex(returns_frame.index)

    else:
        period_inflation = (
            (1.0 + settings.inflation.annual_rate) ** (1.0 / source_periods_per_year) - 1.0
        )
        inflation_series = pd.Series(period_inflation, index=returns_frame.index)

    # Реальная доходность: (1 + номинальная) / (1 + инфляция) − 1.
    returns_frame = (1.0 + returns_frame).div(1.0 + inflation_series, axis=0) - 1.0

    if settings.resample_to and settings.resample_to != settings.frequency:
        if PERIODS_PER_YEAR[settings.resample_to] > source_periods_per_year:
            raise DataError('Можно агрегировать только к более крупному таймфрейму.')

        if not dataset.has_dates:
            raise DataError('Для смены таймфрейма в первом столбце нужны даты.')

        returns_frame = (
            (1.0 + returns_frame).resample(RESAMPLE_RULE[settings.resample_to]).prod() - 1.0
        )

    validate(returns_frame)

    return returns_frame, warnings


def real_rate(nominal_annual: float, settings: Settings) -> float:
    if settings.inflation.source == 'series' and settings.inflation.series:
        unit_divisor = 1.0
        if settings.units == 'percent':
            unit_divisor = 100.0

        period_inflation = float(np.mean(
            np.asarray(settings.inflation.series, dtype=float) / unit_divisor
        ))
        annual_inflation = (1.0 + period_inflation) ** PERIODS_PER_YEAR[settings.frequency] - 1.0
    else:
        annual_inflation = settings.inflation.annual_rate

    return (1.0 + nominal_annual) / (1.0 + annual_inflation) - 1.0


def validate(returns_frame: pd.DataFrame) -> None:
    if not MIN_ASSETS <= returns_frame.shape[1] <= MAX_ASSETS:
        raise DataError(
            f'Нужно от {MIN_ASSETS} до {MAX_ASSETS} активов, получено {returns_frame.shape[1]}.'
        )

    if len(returns_frame) < MIN_OBS:
        raise DataError(
            f'Нужно не менее {MIN_OBS} наблюдений, '
            f'после обработки осталось {len(returns_frame)}.'
        )

    if returns_frame.isna().any().any() or not np.isfinite(returns_frame.values).all():
        raise DataError('В данных остались NaN или бесконечные значения.')

    zero_variance = returns_frame.std() == 0
    if zero_variance.any():
        invalid_assets = ', '.join(returns_frame.columns[zero_variance])
        raise DataError(f'У индикаторов нулевая дисперсия: {invalid_assets}.')


def load_tickers(
    tickers: list[str],
    start: str,
    end: str | None,
    frequency: str,
) -> tuple[Dataset, list[str]]:
    warnings = []
    if frequency not in PERIODS_PER_YEAR:
        raise DataError(f'Неизвестная частота данных: {frequency}.')
    try:
        start_date = pd.Timestamp(start)
        end_date = pd.Timestamp(end) if end else pd.Timestamp.now().normalize()
        if pd.isna(start_date) or pd.isna(end_date) or start_date.tz or end_date.tz:
            raise ValueError('Invalid date')
    except (TypeError, ValueError) as error:
        raise DataError('Укажите корректные даты начала и окончания.') from error
    if start_date >= end_date:
        raise DataError('Дата начала должна быть раньше даты окончания.')

    tickers = [ticker.strip().upper() for ticker in tickers if ticker.strip()]
    if not MIN_ASSETS <= len(tickers) <= MAX_ASSETS:
        raise DataError(f'Нужно от {MIN_ASSETS} до {MAX_ASSETS} рыночных тикеров.')
    if len(set(tickers)) != len(tickers):
        raise DataError('В списке есть повторяющиеся тикеры.')
    if any(EXCLUDED_INDICATOR_RE.search(ticker) for ticker in tickers):
        raise DataError('Показатель M2 исключён. Укажите только рыночные активы.')
    if any(not re.fullmatch(r'(?:MOEX:)?[A-Z0-9^][A-Z0-9^=._-]*', t) for t in tickers):
        raise DataError('Некорректный тикер. Примеры: MOEX:SBER, SPY, BTC-USD.')

    def load_one(ticker):
        if ticker.startswith('MOEX:'):
            # Единая семантика end: правая граница не включается, как в Yahoo.
            till = (end_date - pd.Timedelta(days=1)).date().isoformat()
            series = _load_moex_prices(ticker[5:], start_date.date().isoformat(), till)
        else:
            series = _load_yahoo_prices(ticker, start_date.date().isoformat(), end_date.date().isoformat())
        series = series.copy()
        series.index = pd.to_datetime(series.index).tz_localize(None).normalize()
        series = series[~series.index.duplicated(keep='last')].sort_index()
        series = series.loc[(series.index >= start_date) & (series.index < end_date)]
        if series.empty:
            raise DataError(f'Нет котировок {ticker} за выбранный период.')
        return series

    # MOEX загружается параллельно. Отдельные вызовы yfinance.download идут
    # последовательно: старые поддерживаемые версии используют общее состояние.
    moex_tickers = [ticker for ticker in tickers if ticker.startswith('MOEX:')]
    asset_prices = {}
    if moex_tickers:
        with ThreadPoolExecutor(max_workers=min(4, len(moex_tickers))) as pool:
            asset_prices.update(zip(moex_tickers, pool.map(load_one, moex_tickers)))
    for ticker in tickers:
        if ticker not in asset_prices:
            asset_prices[ticker] = load_one(ticker)
    asset_prices = {ticker: asset_prices[ticker] for ticker in tickers}

    prices_frame = pd.DataFrame(asset_prices).sort_index()
    resample_rule = RESAMPLE_RULE.get(frequency)
    if resample_rule:
        prices_frame = prices_frame.resample(resample_rule).last()
        # Неполный текущий месяц/квартал/год не считается полным наблюдением.
        prices_frame = prices_frame.loc[prices_frame.index < end_date]
    else:
        prices_frame = prices_frame.dropna()
    # Не переносим старые цены в пропуски и не создаём фиктивные нулевые доходности.
    returns_frame = prices_frame.pct_change(fill_method=None).dropna() * 100.0
    if returns_frame.empty:
        raise DataError('Нет общей истории котировок. Измените период или список активов.')
    dataset = Dataset(
        columns=list(returns_frame.columns),
        index=[period.date().isoformat() for period in returns_frame.index],
        data=returns_frame.values.tolist(),
        has_dates=True,
    )
    warnings.append('Загружены исторические рыночные котировки; доходности рассчитаны по ценам закрытия.')
    if any(t.startswith('MOEX:') for t in tickers):
        warnings.append('MOEX: цены CLOSE не скорректированы на дивиденды и дробления акций.')
    if len(dataset.data) < MIN_OBS:
        warnings.append(f'Получено {len(dataset.data)} наблюдений — для анализа нужно не менее {MIN_OBS}.')
    return dataset, warnings


def _load_yahoo_prices(ticker: str, start: str, end: str | None) -> pd.Series:
    import yfinance

    try:
        prices_frame = yfinance.download(
            ticker,
            start=start,
            end=end,
            auto_adjust=True,
            progress=False,
            timeout=20,
        )
    except Exception as error:
        logger.warning('Yahoo Finance request failed for %s: %s', ticker, error)
        raise DataError(f'Не удалось загрузить котировки {ticker} из Yahoo Finance.') from error

    if prices_frame is None or prices_frame.empty or 'Close' not in prices_frame:
        raise DataError(f'Yahoo Finance не вернул цены закрытия по {ticker}.')

    closing_prices = prices_frame['Close']

    if isinstance(closing_prices, pd.DataFrame):
        closing_prices = closing_prices.iloc[:, 0]

    closing_prices = pd.to_numeric(closing_prices, errors='coerce').dropna()
    closing_prices = closing_prices[closing_prices > 0]

    if closing_prices.empty or not np.isfinite(closing_prices.to_numpy(dtype=float)).all():
        raise DataError(f'В данных Yahoo Finance по {ticker} нет корректных цен.')

    return closing_prices.rename(ticker)


def _load_moex_prices(secid: str, start: str, end: str | None) -> pd.Series:
    import requests

    engine_market = 'stock/shares'
    board_path = '/boards/TQBR'

    if secid.upper() in {'IMOEX', 'MCFTR', 'RUCBITR', 'RUCBTRNS', 'RTSI', 'MOEXBC', 'RGBI', 'RGBITR'}:
        engine_market = 'stock/index'
        board_path = ''

    engine_name, market_name = engine_market.split('/')
    url = (
        f'https://iss.moex.com/iss/history/engines/{engine_name}/markets/'
        f'{market_name}{board_path}/securities/{secid}.json'
    )
    closing_prices = {}
    cursor = 0

    while True:
        params = {
            'from': start,
            'start': cursor,
            'iss.meta': 'off',
            'iss.only': 'history',
            'history.columns': 'TRADEDATE,CLOSE',
        }
        if end:
            params['till'] = end

        try:
            response = requests.get(url, params=params, timeout=20)
            response.raise_for_status()
            payload = response.json()
            history = payload.get('history', {})
            rows = history.get('data', [])
            columns = history.get('columns', ['TRADEDATE', 'CLOSE'])
            date_column = columns.index('TRADEDATE')
            close_column = columns.index('CLOSE')
        except (requests.RequestException, ValueError, AttributeError) as error:
            logger.warning('MOEX ISS request failed for %s: %s', secid, error)
            raise DataError(f'Не удалось загрузить котировки {secid} с MOEX.') from error

        if not rows:
            break

        for row in rows:
            if not isinstance(row, (list, tuple)) or len(row) <= max(date_column, close_column) or row[close_column] is None:
                continue

            try:
                price = float(row[close_column])
                if np.isfinite(price) and price > 0:
                    closing_prices[pd.Timestamp(row[date_column])] = price
            except (TypeError, ValueError):
                continue

        cursor += len(rows)

    if not closing_prices:
        raise DataError(f'MOEX ISS не вернул данных по {secid}.')

    return pd.Series(closing_prices).sort_index().rename(f'MOEX:{secid}')
